"""
Shared visual design system for the .xlsx report: brand colors, fonts, and
small style helpers, kept in one place so every tab in report.py looks like
it came from the same product rather than five independently-styled sheets.

Palette is charcoal + amber, deliberately not the dark-navy-blue "AI
dashboard" default -- charcoal reads as serious/technical without leaning on
a color every other generated report already uses, and amber gives the
banner/section-header accents a single recognizable signature instead of
just more blue.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from adforensictool.version import __version__ as TOOL_VERSION  # noqa: F401  (re-exported for report.py)

TOOL_NAME = "ADForensicTool"

# -- Brand palette ----------------------------------------------------------
BRAND_DARK = "232323"    # primary brand color: banners, tab headers
BRAND_DARK_ALT = "1A1A1A"  # slightly deeper shade, used sparingly for depth
ACCENT_AMBER = "D9A521"  # signature accent: banner divider, section underlines
SLATE = "6B6B6B"         # muted secondary text (neutral gray, no blue tint)
LIGHT_GRAY = "F2F2F2"    # zebra striping
MID_GRAY = "D9D9D9"      # borders/dividers
WHITE = "FFFFFF"

# -- Severity palette (Office accent colors) --------------------------------
# Plain hex values (no ARGB prefix) -- the single source of truth for
# severity color, reused for cell fills below AND for chart data-point
# colors, so a bar chart bucket for "Critical" is guaranteed to be the same
# red as a "Critical" table cell rather than an arbitrary Excel theme color.
SEVERITY_HEX = {
    "critical": "C00000",
    "high": "ED7D31",
    "medium": "FFC000",
    "low": "70AD47",
}

SEVERITY_FILLS = {
    sev: PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
    for sev, hex_color in SEVERITY_HEX.items()
}
SEVERITY_ROW_FILLS = {
    "critical": PatternFill(start_color="FBE0E0", end_color="FBE0E0", fill_type="solid"),
    "high": PatternFill(start_color="FCE7D6", end_color="FCE7D6", fill_type="solid"),
    "medium": PatternFill(start_color="FFF4CC", end_color="FFF4CC", fill_type="solid"),
    "low": PatternFill(start_color="E5F0E1", end_color="E5F0E1", fill_type="solid"),
}
SEVERITY_FONTS = {
    "critical": Font(color=WHITE, bold=True),
    "high": Font(color=WHITE, bold=True),
    "medium": Font(color="7F5F00", bold=True),
    "low": Font(color=WHITE, bold=True),
}
SEVERITY_TEXT_COLORS = {
    "critical": "C00000",
    "high": "C55A11",
    "medium": "7F5F00",
    "low": "548235",
}
SEVERITY_ORDER = ["critical", "high", "medium", "low"]

# -- Reusable styles ----------------------------------------------------------
BANNER_FILL = PatternFill(start_color=BRAND_DARK, end_color=BRAND_DARK, fill_type="solid")
BANNER_ACCENT_FILL = PatternFill(start_color=ACCENT_AMBER, end_color=ACCENT_AMBER, fill_type="solid")
BANNER_TITLE_FONT = Font(color=WHITE, bold=True, size=20, name="Calibri")
BANNER_SUBTITLE_FONT = Font(color="B8B8B8", size=10, name="Calibri")
BANNER_VERSION_FONT = Font(color=ACCENT_AMBER, bold=True, size=10, name="Calibri")

SECTION_HEADER_FONT = Font(color=BRAND_DARK, bold=True, size=12, name="Calibri")
SECTION_HEADER_BORDER = Border(bottom=Side(style="medium", color=ACCENT_AMBER))

TABLE_HEADER_FILL = PatternFill(start_color=BRAND_DARK, end_color=BRAND_DARK, fill_type="solid")
TABLE_HEADER_FONT = Font(color=WHITE, bold=True, name="Calibri")
TABLE_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

ZEBRA_FILL = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color=MID_GRAY),
    right=Side(style="thin", color=MID_GRAY),
    top=Side(style="thin", color=MID_GRAY),
    bottom=Side(style="thin", color=MID_GRAY),
)

KPI_LABEL_FONT = Font(color=SLATE, size=9, name="Calibri")
KPI_VALUE_FONT = Font(color=BRAND_DARK, bold=True, size=22, name="Calibri")
KPI_FILL = PatternFill(start_color="F7F7F5", end_color="F7F7F5", fill_type="solid")
KPI_BORDER = Border(
    left=Side(style="thin", color=MID_GRAY),
    right=Side(style="thin", color=MID_GRAY),
    top=Side(style="thin", color=MID_GRAY),
    bottom=Side(style="thin", color=MID_GRAY),
)


def risk_band(score: int) -> tuple[str, str]:
    """Return (label, hex_color) for a 0-100 risk score."""
    if score >= 80:
        return "CRITICAL", "C00000"
    if score >= 60:
        return "HIGH", "ED7D31"
    if score >= 30:
        return "ELEVATED", "FFC000"
    return "LOW", "70AD47"
