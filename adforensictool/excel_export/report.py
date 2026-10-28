"""
Builds the analyst-facing .xlsx report from flagged events.

Tabs:
  - Executive Summary       branded dashboard: risk score, KPI tiles, charts,
                            attack narrative (one sentence per chain), key findings
  - Attack Chains           one row per correlated chain + a stage-by-stage breakdown
  - Detections by Rule      flagged events grouped per rule / threshold cluster
  - Raw Flagged Events      full detail table, color-coded by severity, filterable
  - Account/Group Changes   dedicated view of privilege-escalation events
  - MITRE ATT&CK Coverage   technique -> observed count heatmap + chart

Timestamps are written as real Excel datetimes (sortable, filterable) in the
display timezone, which is named in every timestamp column header and on
the Executive Summary -- Excel itself has no notion of timezones.

Visual design (colors, fonts, KPI tile styling) lives in excel_export/theme.py
so every tab shares one look rather than independently-styled sheets.
"""
from __future__ import annotations

import math
from datetime import UTC, datetime, tzinfo
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from adforensictool.correlator.chain_builder import Chain, build_chains
from adforensictool.excel_export import theme
from adforensictool.narrator.narrate import chain_headline, narrate_chain
from adforensictool.parser.evtx_parser import ParseStats
from adforensictool.rules.engine import Rule

SEVERITY_SCORE = {"critical": 40, "high": 20, "medium": 10, "low": 5}

# Human-readable names for the MITRE technique IDs this tool's default rules
# reference. Extend when you add rules that cite new technique IDs.
MITRE_NAMES = {
    "T1110": "Brute Force",
    "T1110.003": "Brute Force: Password Spraying",
    "T1078": "Valid Accounts",
    "T1098": "Account Manipulation",
    "T1070.001": "Clear Windows Event Logs",
    "T1558.003": "Steal or Forge Kerberos Tickets: Kerberoasting",
    "T1558.004": "Steal or Forge Kerberos Tickets: AS-REP Roasting",
    "T1003.006": "OS Credential Dumping: DCSync",
    "T1550.002": "Use Alternate Authentication Material: Pass the Hash",
    "T1136.002": "Create Account: Domain Account",
    "T1543": "Create or Modify System Process",
    "T1543.003": "Create or Modify System Process: Windows Service",
    "T1053.005": "Scheduled Task/Job: Scheduled Task",
    "T1562.002": "Impair Defenses: Disable Windows Event Logging",
}

DISPLAY_COLUMNS = [
    "timestamp",
    "event_id",
    "record_id",
    "rule_name",
    "severity",
    "mitre_technique",
    "chain_id",
    "user",
    "user_sid",
    "subject_user",
    "target_group",
    "source_ip",
    "target_host",
    "logon_type",
    "process_name",
    "raw_message",
    "cluster_id",
]
OPTIONAL_DISPLAY_COLUMNS = ["source_file"]  # shown only when present (multi-file input)

DASHBOARD_LAST_COL = 12  # column L -- width of the banner/KPI/bullet area
MAX_NARRATED_CHAINS = 10  # on the Executive Summary; the rest are on the Attack Chains tab
TOP_RULES_IN_CHART = 8
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"


def zone_label(tz: tzinfo | None) -> str:
    if tz is None:
        return "UTC"
    return getattr(tz, "key", None) or str(tz)


def _local_naive(value: Any, tz: tzinfo | None) -> Any:
    """tz-aware timestamp -> naive datetime in the display zone (what Excel
    can store); other values pass through. NaT -> empty."""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    if isinstance(value, pd.Timestamp | datetime):
        ts = pd.Timestamp(value)
        if ts.tzinfo is not None:
            ts = ts.tz_convert(tz or UTC).tz_localize(None)
        return ts.to_pydatetime()
    return value


def _fmt_time(value: Any, tz: tzinfo | None) -> str:
    local = _local_naive(value, tz)
    return local.strftime("%Y-%m-%d %H:%M:%S") if isinstance(local, datetime) else "?"


# --------------------------------------------------------------------------
# Generic table helpers, shared by the data tabs
# --------------------------------------------------------------------------


def _set_value(cell, value: Any) -> None:
    """Assign value to cell, never letting log-derived text become a formula.

    Usernames, hostnames, service names etc. are attacker-controlled: anyone
    can trigger a failed logon for an account literally named
    '=HYPERLINK("http://evil/",...)'. openpyxl stores any string starting
    with "=" as a live formula, so force such values back to plain text.
    The analyst still sees the exact original string.
    """
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    elif isinstance(value, datetime):
        cell.number_format = DATETIME_FORMAT


def _write_cell(ws: Worksheet, row: int, column: int, value: Any):
    cell = ws.cell(row=row, column=column)
    _set_value(cell, value)
    return cell


def _header_label(column: str, tz: tzinfo | None) -> str:
    if column in {"timestamp", "start", "end", "first_seen", "last_seen"}:
        return f"{column} ({zone_label(tz)})"
    return column


def _write_header(ws: Worksheet, columns: list[str], row: int = 1, tz: tzinfo | None = None) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        cell = _write_cell(ws, row, col_idx, _header_label(col_name, tz))
        cell.fill = theme.TABLE_HEADER_FILL
        cell.font = theme.TABLE_HEADER_FONT
        cell.alignment = theme.TABLE_HEADER_ALIGN
        cell.border = theme.THIN_BORDER
    ws.row_dimensions[row].height = 20


def _autosize_columns(ws: Worksheet, columns: list[str], rows: list[list[Any]], tz: tzinfo | None = None) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(_header_label(col_name, tz))
        for row in rows:
            val = row[col_idx - 1] if col_idx - 1 < len(row) else ""
            max_len = max(max_len, 19 if isinstance(val, datetime) else len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)


def _write_rows(
    ws: Worksheet,
    columns: list[str],
    rows: list[list[Any]],
    start_row: int = 2,
    tz: tzinfo | None = None,
    zebra: bool = True,
) -> None:
    """Write data rows under a header, colored by their `severity` column."""
    severity_idx = columns.index("severity") if "severity" in columns else None
    for offset, row in enumerate(rows):
        row_idx = start_row + offset
        for col_idx, value in enumerate(row, start=1):
            cell = _write_cell(ws, row_idx, col_idx, _local_naive(value, tz))
            cell.border = theme.THIN_BORDER

        severity = str(row[severity_idx]).lower() if severity_idx is not None else None
        row_fill = theme.SEVERITY_ROW_FILLS.get(severity) if severity else None
        text_color = theme.SEVERITY_TEXT_COLORS.get(severity) if severity else None
        if row_fill is None and zebra:
            # No severity color applies (or this table doesn't carry one) --
            # fall back to plain zebra striping so the table still reads
            # cleanly at a glance instead of a flat wall of white rows.
            row_fill = theme.ZEBRA_FILL if row_idx % 2 == 0 else None

        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if row_fill:
                cell.fill = row_fill
            if severity_idx == col_idx - 1 and text_color:
                cell.font = Font(color=text_color, bold=True)


def _write_table(ws: Worksheet, df: pd.DataFrame, columns: list[str], tz: tzinfo | None = None) -> None:
    _write_header(ws, columns, tz=tz)
    rows = df[columns].astype(object).where(pd.notnull(df[columns]), "").values.tolist() if not df.empty else []
    _write_rows(ws, columns, rows, tz=tz)
    _autosize_columns(ws, columns, rows, tz)
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    ws.freeze_panes = "A2"


# --------------------------------------------------------------------------
# Executive Summary: branded dashboard
# --------------------------------------------------------------------------


def _fill_range(ws: Worksheet, cell_range: str, fill) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = fill


def _build_banner(ws: Worksheet, source_file: str | None, generated_at: datetime) -> None:
    last_col_letter = get_column_letter(DASHBOARD_LAST_COL)

    _fill_range(ws, f"A1:{last_col_letter}3", theme.BANNER_FILL)

    ws.merge_cells("A1:F2")
    title_cell = ws["A1"]
    title_cell.value = theme.TOOL_NAME
    title_cell.font = theme.BANNER_TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells(f"G1:{last_col_letter}1")
    gen_cell = ws["G1"]
    gen_cell.value = f"Generated: {generated_at.astimezone(UTC).strftime('%Y-%m-%d %H:%M UTC')}"
    gen_cell.font = theme.BANNER_SUBTITLE_FONT
    gen_cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells(f"G2:{last_col_letter}2")
    src_cell = ws["G2"]
    _set_value(src_cell, f"Source: {source_file or 'unknown'}")
    src_cell.font = theme.BANNER_SUBTITLE_FONT
    src_cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells(f"A3:{last_col_letter}3")
    subtitle_cell = ws["A3"]
    # Two-tone rich text: gray subtitle, amber version tag -- the signature
    # accent color, not just another wash of the banner's charcoal.
    subtitle_cell.value = CellRichText(
        TextBlock(InlineFont(color="B8B8B8", sz=10), "Active Directory Security Incident Report   "),
        TextBlock(InlineFont(color=theme.ACCENT_AMBER, sz=10, b=True), f"v{theme.TOOL_VERSION}"),
    )
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Amber divider strip under the charcoal banner -- the recurring accent
    # motif, rather than yet more blue.
    _fill_range(ws, f"A4:{last_col_letter}4", theme.BANNER_ACCENT_FILL)

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 5


def _write_kpi_tile(ws: Worksheet, start_col: int, label: str, value: str, value_color: str) -> None:
    end_col = start_col + 1
    start_letter = get_column_letter(start_col)
    end_letter = get_column_letter(end_col)

    ws.merge_cells(f"{start_letter}6:{end_letter}6")
    label_cell = ws[f"{start_letter}6"]
    label_cell.value = label
    label_cell.font = theme.KPI_LABEL_FONT
    label_cell.alignment = Alignment(horizontal="center", vertical="bottom")

    ws.merge_cells(f"{start_letter}7:{end_letter}8")
    value_cell = ws[f"{start_letter}7"]
    value_cell.value = value
    value_cell.font = Font(color=value_color, bold=True, size=22, name="Calibri")
    value_cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws[f"{start_letter}6:{end_letter}8"]:
        for cell in row:
            cell.fill = theme.KPI_FILL
            cell.border = theme.KPI_BORDER


def _build_kpi_row(
    ws: Worksheet,
    records_read: int,
    flagged: pd.DataFrame,
    chains: list[Chain],
    risk_score: int,
    severity_counts: dict,
) -> None:
    risk_label, risk_color = theme.risk_band(risk_score)
    _write_kpi_tile(ws, 2, f"OVERALL RISK ({risk_label})", f"{risk_score}/100", risk_color)
    _write_kpi_tile(ws, 4, "LOG RECORDS READ", str(records_read), theme.BRAND_DARK)
    _write_kpi_tile(ws, 6, "EVENTS FLAGGED", str(len(flagged)), theme.BRAND_DARK)
    chain_color = theme.SEVERITY_TEXT_COLORS["critical"] if chains else theme.BRAND_DARK
    _write_kpi_tile(ws, 8, "ATTACK CHAINS", str(len(chains)), chain_color)
    _write_kpi_tile(ws, 10, "CRITICAL FINDINGS", str(severity_counts.get("critical", 0)), theme.SEVERITY_TEXT_COLORS["critical"])


def _show_axes(chart: BarChart) -> None:
    """openpyxl 3.1 writes chart axes as deleted unless told otherwise, so
    Excel draws bars with no category labels or scale -- a bar chart of
    rule names without the rule names."""
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def _colored_data_points(hex_colors: list[str]) -> list[DataPoint]:
    """Explicit per-bar fill colors for a single-series bar chart.

    Without this, Excel falls back to its default theme accent palette
    (blue, then orange, then grey, ...) assigned by category *position*,
    which is meaningless here -- "Critical" would render blue just because
    it happened to be the first bar. This pins each bar to the same
    severity color used everywhere else in the report (theme.SEVERITY_HEX).
    """
    return [DataPoint(idx=i, spPr=GraphicalProperties(solidFill=color)) for i, color in enumerate(hex_colors)]


def _build_charts(ws: Worksheet, flagged: pd.DataFrame) -> int:
    """Writes hidden chart-data staging tables off to the side, builds two
    native bar charts from them, and returns the row number where content
    below the charts should resume.
    """
    if flagged.empty:
        return 11

    # -- Severity breakdown staging table (columns N:O), hidden from view --
    severity_counts = flagged["severity"].str.lower().value_counts()
    sev_rows = [
        (sev, sev.title(), int(severity_counts.get(sev, 0)))
        for sev in theme.SEVERITY_ORDER
        if severity_counts.get(sev, 0)
    ]

    ws["N1"] = "Severity"
    ws["O1"] = "Count"
    for i, (_, sev_label, count) in enumerate(sev_rows, start=2):
        _write_cell(ws, i, 14, sev_label)
        _write_cell(ws, i, 15, count)

    severity_chart = BarChart()
    severity_chart.type = "col"
    severity_chart.title = "Findings by Severity"
    severity_chart.y_axis.title = "Events"
    severity_chart.x_axis.title = "Severity"
    severity_chart.height = 8
    severity_chart.width = 14
    data = Reference(ws, min_col=15, min_row=1, max_row=1 + len(sev_rows))
    cats = Reference(ws, min_col=14, min_row=2, max_row=1 + len(sev_rows))
    severity_chart.add_data(data, titles_from_data=True)
    severity_chart.set_categories(cats)
    severity_chart.legend = None
    severity_chart.varyColors = True
    severity_chart.series[0].data_points = _colored_data_points([theme.SEVERITY_HEX[sev] for sev, _, _ in sev_rows])
    # Source data lives in columns N:O, which are hidden below -- charts
    # default to plotting only visible cells, so without this the chart
    # renders as a blank box once its data columns are hidden.
    severity_chart.visible_cells_only = False
    _show_axes(severity_chart)
    ws.add_chart(severity_chart, "B11")

    # -- Findings-by-rule staging table (columns Q:R), hidden from view --
    # Keep each rule's severity alongside its name/count so its bar can be
    # colored the same way (one rule always has exactly one severity).
    # Only the top rules fit legibly in the chart (with more, Excel silently
    # drops every other category label); the full list is under KEY FINDINGS.
    # Horizontal bar charts draw the first category at the bottom, so sort
    # ascending to put the largest bar on top.
    by_rule = (
        flagged.groupby(["rule_name", "severity"]).size().reset_index(name="count")
        .sort_values("count", ascending=False, kind="stable").head(TOP_RULES_IN_CHART)
        .iloc[::-1]
    )  # fmt: skip

    ws["Q1"] = "Rule"
    ws["R1"] = "Count"
    for i, row in enumerate(by_rule.itertuples(index=False), start=2):
        _write_cell(ws, i, 17, row.rule_name)
        _write_cell(ws, i, 18, int(row.count))

    rule_chart = BarChart()
    rule_chart.type = "bar"  # horizontal bars read better with long rule names
    rule_chart.title = "Findings by Rule" if len(by_rule) < TOP_RULES_IN_CHART else f"Top {TOP_RULES_IN_CHART} Rules by Findings"
    rule_chart.y_axis.title = "Events"  # in a horizontal bar chart x_axis is the category (rule name) axis
    rule_chart.height = 8
    rule_chart.width = 14
    data = Reference(ws, min_col=18, min_row=1, max_row=1 + len(by_rule))
    cats = Reference(ws, min_col=17, min_row=2, max_row=1 + len(by_rule))
    rule_chart.add_data(data, titles_from_data=True)
    rule_chart.set_categories(cats)
    rule_chart.legend = None
    rule_chart.varyColors = True
    rule_chart.series[0].data_points = _colored_data_points(
        [theme.SEVERITY_HEX.get(str(sev).lower(), theme.SLATE) for sev in by_rule["severity"]]
    )
    rule_chart.visible_cells_only = False  # same hidden-source-columns issue as above
    _show_axes(rule_chart)
    ws.add_chart(rule_chart, "H11")

    for col in ("N", "O", "Q", "R"):
        ws.column_dimensions[col].hidden = True

    return 28  # generous clearance below ~8cm-tall charts anchored at row 11


def _build_section_header(ws: Worksheet, row: int, text: str) -> None:
    last_col_letter = get_column_letter(DASHBOARD_LAST_COL)
    ws.merge_cells(f"A{row}:{last_col_letter}{row}")
    cell = ws[f"A{row}"]
    _set_value(cell, text)
    cell.font = theme.SECTION_HEADER_FONT
    for c in ws[f"A{row}:{last_col_letter}{row}"][0]:
        c.border = theme.SECTION_HEADER_BORDER


def _write_wide_row(ws: Worksheet, row: int, text: str, font: Font | None = None) -> None:
    """A full-width merged text row. Merged cells never auto-fit their
    height in Excel, so estimate the wrapped line count from the text."""
    last_col_letter = get_column_letter(DASHBOARD_LAST_COL)
    ws.merge_cells(f"A{row}:{last_col_letter}{row}")
    cell = ws[f"A{row}"]
    _set_value(cell, text)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    if font:
        cell.font = font
    lines = max(1, math.ceil(len(text) / 130))  # ~130 chars fit across the 12 dashboard columns
    ws.row_dimensions[row].height = 18 if lines == 1 else 15 * lines + 4


def _build_narrative(ws: Worksheet, r: int, chains: list[Chain], tz: tzinfo | None) -> int:
    _build_section_header(ws, r, "ATTACK NARRATIVE")
    r += 2
    for chain in chains[:MAX_NARRATED_CHAINS]:
        color = theme.SEVERITY_TEXT_COLORS.get(chain.severity, theme.BRAND_DARK)
        _write_wide_row(ws, r, f"  ■  [{chain.severity.upper()}] {chain.id}: {chain_headline(chain)}", Font(color=color, bold=True))
        r += 1
        _write_wide_row(ws, r, f"      {narrate_chain(chain, tz)}")
        r += 2
    if len(chains) > MAX_NARRATED_CHAINS:
        _write_wide_row(
            ws,
            r,
            f"  + {len(chains) - MAX_NARRATED_CHAINS} more chain(s) on the Attack Chains tab.",
            Font(italic=True, color=theme.SLATE),
        )
        r += 2
    return r


def _build_executive_summary(
    ws: Worksheet,
    df: pd.DataFrame,
    flagged: pd.DataFrame,
    chains: list[Chain],
    source_file: str | None,
    generated_at: datetime,
    parse_stats: ParseStats | None = None,
    suppressed: dict[str, int] | None = None,
    tz: tzinfo | None = None,
) -> None:
    for col in range(1, DASHBOARD_LAST_COL + 1):
        ws.column_dimensions[get_column_letter(col)].width = 11

    severity_counts = flagged["severity"].str.lower().value_counts().to_dict() if not flagged.empty else {}
    risk_score = min(100, sum(SEVERITY_SCORE.get(sev, 0) * count for sev, count in severity_counts.items()))

    _build_banner(ws, source_file, generated_at)
    # With the Event-ID prefilter, df holds only rule-relevant events; the
    # number of records actually read comes from the parse stats.
    records_read = parse_stats.records_read if parse_stats is not None and parse_stats.records_read else len(df)
    _build_kpi_row(ws, records_read, flagged, chains, risk_score, severity_counts)
    r = _build_charts(ws, flagged)

    if chains:
        r = _build_narrative(ws, r, chains, tz)

    _build_section_header(ws, r, "KEY FINDINGS")
    r += 2

    if flagged.empty:
        _write_wide_row(ws, r, "No detections fired against the supplied log.")
        r += 1
    else:
        by_rule = flagged.groupby(["rule_id", "rule_name", "severity"]).agg(
            count=("event_id", "size"),
            first_seen=("timestamp", "min"),
            last_seen=("timestamp", "max"),
        ).reset_index()
        by_rule = by_rule.sort_values("first_seen")
        for _, row in by_rule.iterrows():
            severity = str(row["severity"]).lower()
            first, last = _fmt_time(row["first_seen"], tz), _fmt_time(row["last_seen"], tz)
            when = f"at {first}" if first == last else f"between {first} and {last}"
            bullet = f"  ●  [{row['severity'].upper()}] {row['rule_name']}: {row['count']} event(s) {when}"
            font = Font(color=theme.SEVERITY_TEXT_COLORS.get(severity, theme.BRAND_DARK), bold=True)
            _write_wide_row(ws, r, bullet, font)
            r += 1

    if parse_stats is not None and parse_stats.records_skipped:
        r += 1
        _write_wide_row(
            ws,
            r,
            (
                f"  ▲  {parse_stats.records_skipped} of {parse_stats.records_read} log record(s) could not be "
                "parsed and were skipped. A corrupted export can cause this, but so can tampering: "
                "verify the source log's integrity."
            ),
            Font(color=theme.SEVERITY_TEXT_COLORS["high"], bold=True),
        )
        r += 1

    for file, channel, computer, low, high, missing in parse_stats.record_id_gaps() if parse_stats is not None else []:
        r += 1
        name = file.replace("\\", "/").rsplit("/", 1)[-1]
        _write_wide_row(
            ws,
            r,
            (
                f"  ▲  {name} [{channel or '?'} log, {computer or '?'}]: {missing} EventRecordID(s) missing between "
                f"{low} and {high}. Expected if this is a filtered export; in a complete export it means records "
                "were lost or individually deleted (an anti-forensics technique that leaves no 1102 event)."
            ),
            Font(color=theme.SEVERITY_TEXT_COLORS["high"], bold=True),
        )
        r += 1

    if parse_stats is not None and parse_stats.records_duplicate:
        r += 1
        _write_wide_row(
            ws,
            r,
            f"  ○  {parse_stats.records_duplicate} record(s) appeared in more than one input file and were counted once.",
            Font(italic=True, color=theme.SLATE),
        )
        r += 1

    if suppressed:
        r += 1
        total = sum(suppressed.values())
        detail = "; ".join(f"{desc} ({count})" for desc, count in suppressed.items())
        _write_wide_row(
            ws,
            r,
            f"  ○  {total} flagged event(s) suppressed by the allowlist and not shown in this report: {detail}",
            Font(italic=True, color=theme.SLATE),
        )
        r += 1

    r += 1
    _write_wide_row(ws, r, f"All event times in this report are shown in {zone_label(tz)}.", Font(italic=True, color=theme.SLATE))


# --------------------------------------------------------------------------
# Attack Chains / Detections / Account-Group Changes / MITRE Coverage tabs
# --------------------------------------------------------------------------

CHAIN_COLUMNS = ["chain_id", "severity", "start", "end", "accounts", "source_ips", "tactics", "stages", "event_count", "narrative"]
STAGE_COLUMNS = ["chain_id", "stage", "start", "end", "rule_name", "severity", "tactic", "mitre_technique", "event_count", "users", "source_ips"]


def _build_attack_chains(ws: Worksheet, chains: list[Chain], tz: tzinfo | None) -> None:
    _write_header(ws, CHAIN_COLUMNS, tz=tz)
    if not chains:
        _write_cell(ws, 2, 1, "No multi-stage attack chains were identified: no two different detections shared an account within the correlation window.")
        _autosize_columns(ws, CHAIN_COLUMNS, [], tz)
        ws.column_dimensions["A"].width = 30
        return

    chain_rows = [
        [
            c.id,
            c.severity,
            c.start_time,
            c.end_time,
            ", ".join(c.users),
            ", ".join(c.source_ips),
            " → ".join(c.tactics),
            " → ".join(dict.fromkeys(s.rule_name for s in c.stages)),
            c.event_count,
            narrate_chain(c, tz),
        ]
        for c in chains
    ]
    _write_rows(ws, CHAIN_COLUMNS, chain_rows, tz=tz)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CHAIN_COLUMNS))}{len(chain_rows) + 1}"

    # Stage-by-stage breakdown below the chain table.
    r = len(chain_rows) + 3
    _write_cell(ws, r, 1, "Chain stages").font = theme.SECTION_HEADER_FONT
    r += 1
    _write_header(ws, STAGE_COLUMNS, row=r, tz=tz)
    stage_rows = [
        [c.id, n, s.start_time, s.end_time, s.rule_name, s.severity, s.tactic, s.mitre_technique, s.event_count, ", ".join(s.users), ", ".join(s.source_ips)]
        for c in chains
        for n, s in enumerate(c.stages, start=1)
    ]
    _write_rows(ws, STAGE_COLUMNS, stage_rows, start_row=r + 1, tz=tz)

    # Column widths: fit both tables; the narrative column wraps instead.
    widths: dict[int, int] = {}
    for columns, rows in ((CHAIN_COLUMNS, chain_rows), (STAGE_COLUMNS, stage_rows)):
        for i, name in enumerate(columns, start=1):
            longest = max([len(_header_label(name, tz))] + [19 if isinstance(row[i - 1], pd.Timestamp) else len(str(row[i - 1])) for row in rows])
            widths[i] = max(widths.get(i, 0), min(longest + 2, 50))
    for i, width in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = width
    # Long free-text columns wrap at a fixed width; every cell in a chain row
    # is top-aligned so short values sit level with the start of the text.
    # (accounts shares its column with the stage table's rule_name below)
    wrap_widths = {"accounts": 45, "tactics": 40, "stages": 50, "narrative": 100}
    for name, width in wrap_widths.items():
        ws.column_dimensions[get_column_letter(CHAIN_COLUMNS.index(name) + 1)].width = width
    for row_idx, row in enumerate(chain_rows, start=2):
        lines = 1
        for col_idx, name in enumerate(CHAIN_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(wrap_text=name in wrap_widths, vertical="top")
            if name in wrap_widths:
                lines = max(lines, math.ceil(len(str(row[col_idx - 1])) / (wrap_widths[name] * 0.95)))
        ws.row_dimensions[row_idx].height = 15 * lines + 4
    ws.freeze_panes = "B2"


def _build_detection_groups(ws: Worksheet, flagged: pd.DataFrame, tz: tzinfo | None) -> None:
    columns = ["rule_id", "rule_name", "severity", "mitre_technique", "cluster_id", "event_count", "first_seen", "last_seen"]
    if flagged.empty:
        _write_header(ws, columns, tz=tz)
        _autosize_columns(ws, columns, [], tz)
        return

    group_cols = ["rule_id", "rule_name", "severity", "mitre_technique", "cluster_id"]
    grouped = flagged.groupby(group_cols, dropna=False).agg(
        event_count=("event_id", "size"),
        first_seen=("timestamp", "min"),
        last_seen=("timestamp", "max"),
    ).reset_index()
    grouped = grouped.sort_values("first_seen")
    _write_table(ws, grouped, columns, tz)


def _build_account_changes(ws: Worksheet, flagged: pd.DataFrame, tz: tzinfo | None) -> None:
    columns = ["timestamp", "target_host", "event_id", "user", "target_group", "subject_user", "severity", "raw_message"]
    changes = flagged[flagged["rule_id"] == "privilege_escalation"] if not flagged.empty else flagged
    _write_table(ws, changes, columns, tz)


def _build_mitre_coverage(ws: Worksheet, flagged: pd.DataFrame, rules: list[Rule]) -> None:
    columns = ["mitre_technique", "technique_name", "observed_count"]
    _write_header(ws, columns)

    watched_techniques = sorted({r.mitre_technique for r in rules})
    counts = flagged.groupby("mitre_technique").size().to_dict() if not flagged.empty else {}

    rows = []
    for technique in watched_techniques:
        rows.append([technique, MITRE_NAMES.get(technique, "Unknown"), counts.get(technique, 0)])

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = _write_cell(ws, row_idx, col_idx, value)
            cell.border = theme.THIN_BORDER
            if row_idx % 2 == 0:
                cell.fill = theme.ZEBRA_FILL

    _autosize_columns(ws, columns, rows)

    if rows:
        last_row = len(rows) + 1
        count_col_letter = get_column_letter(columns.index("observed_count") + 1)
        color_scale = ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="C00000")
        ws.conditional_formatting.add(f"{count_col_letter}2:{count_col_letter}{last_row}", color_scale)

        chart_col_start = len(columns) + 3  # a couple columns of gutter to the right of the table
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "MITRE ATT&CK Technique Coverage"
        chart.y_axis.title = "Events Observed"  # value axis (horizontal bars)
        chart.height = max(6, len(rows) * 1.1)
        chart.width = 16
        data = Reference(ws, min_col=3, min_row=1, max_row=last_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        _show_axes(chart)
        ws.add_chart(chart, f"{get_column_letter(chart_col_start)}1")

    ws.freeze_panes = "A2"


def chain_ids_by_row(flagged: pd.DataFrame, chains: list[Chain]) -> pd.Series:
    """flagged-row index -> chain id ("" when the row is in no chain)."""
    mapping = {idx: chain.id for chain in chains for stage in chain.stages for idx in stage._index}
    return pd.Series([mapping.get(idx, "") for idx in flagged.index], index=flagged.index, dtype=object)


def build_report(
    df: pd.DataFrame,
    flagged: pd.DataFrame,
    rules: list[Rule],
    output_path: str,
    source_file: str | None = None,
    generated_at: datetime | None = None,
    parse_stats: ParseStats | None = None,
    suppressed: dict[str, int] | None = None,
    chains: list[Chain] | None = None,
    tz: tzinfo | None = None,
) -> None:
    """Build the .xlsx report and write it to output_path.

    df:            full normalized event DataFrame (event counts, provenance).
    flagged:       output of rules.engine.apply_rules -- one row per flagged event.
    rules:         the loaded rule definitions, used to enumerate MITRE
                   techniques this run *watched for* (not just ones that fired).
    source_file:   input path(s) description, shown in the Executive Summary
                   banner for provenance. Optional -- falls back to "unknown".
    generated_at:  report generation timestamp shown in the banner. Defaults
                   to now (UTC) if not supplied.
    parse_stats:   record read/skip counters from parse_evtx; a non-zero skip
                   count is called out on the Executive Summary.
    suppressed:    {allowlist entry description: events it suppressed}, from
                   rules.allowlist.apply_allowlist; summarized on the
                   Executive Summary so hidden findings are never silent.
    chains:        correlator output for `flagged`; built here with the
                   default window when not supplied.
    tz:            display timezone for every event time (default UTC).
    """
    generated_at = generated_at or datetime.now(UTC)
    if chains is None:
        chains = build_chains(flagged)
    flagged = flagged.copy()
    flagged["chain_id"] = chain_ids_by_row(flagged, chains) if not flagged.empty else pd.Series(dtype=object)
    display_columns = DISPLAY_COLUMNS + [c for c in OPTIONAL_DISPLAY_COLUMNS if c in flagged.columns]

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    _build_executive_summary(ws_summary, df, flagged, chains, source_file, generated_at, parse_stats, suppressed, tz)

    _build_attack_chains(wb.create_sheet("Attack Chains"), chains, tz)
    _build_detection_groups(wb.create_sheet("Detections by Rule"), flagged, tz)
    _write_table(wb.create_sheet("Raw Flagged Events"), flagged.reindex(columns=display_columns), display_columns, tz)
    _build_account_changes(wb.create_sheet("Account-Group Changes"), flagged, tz)
    _build_mitre_coverage(wb.create_sheet("MITRE ATT&CK Coverage"), flagged, rules)

    wb.save(output_path)
